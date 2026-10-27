
import os
for _key in ("OMP_NUM_THREADS","MKL_NUM_THREADS","OPENBLAS_NUM_THREADS"): os.environ.__setitem__(_key,"1")


NEW_PROBLEMS = ['sphere2', 'sphere6', 'rosenbrock2', 'rosenbrock6', 'ackley6']
PROBLEMS = ['branin', 'hartmann3', 'borehole', 'sphere2', 'sphere6', 'rosenbrock2', 'rosenbrock6', 'ackley6']


from dataclasses import dataclass
import numpy as np
from scipy.stats import qmc


def _p001(n, d, seed):
    if n < 1:
        return np.empty((0, d))
    return qmc.Sobol(d, scramble=True, seed=seed).random_base2(int(np.ceil(np.log2(n))))[:n]


@dataclass
class Problem:
    name: str
    discrepancy: float = 1.0
    stress: bool = False

    @property
    def dim(self):
        if self.name not in PROBLEMS:raise ValueError('Unsupported benchmark: '+self.name)
        return {"branin": 2, "hartmann3": 3, "borehole": 8, **{name: int(name[-1]) for name in NEW_PROBLEMS}}[self.name]

    def raw(self, x):
        x = np.atleast_2d(np.asarray(x, dtype=float))
        if x.shape[1] != self.dim or np.any((x < 0) | (x > 1)):
            raise ValueError("Expected points in the unit cube with correct dimension")
        if self.name in NEW_PROBLEMS:
            family = self.name[:-1]
            if family == "sphere":
                z = -5.12 + 10.24*x
                return np.sum(z*z, axis=1)
            if family == "rosenbrock":
                z = -2.048 + 4.096*x
                return np.sum(100*(z[:,1:]-z[:,:-1]**2)**2+(1-z[:,:-1])**2, axis=1)
            if family == "ackley":

                z = -5 + 10*x
                return -20*np.exp(-.2*np.sqrt(np.mean(z*z,axis=1)))-np.exp(np.mean(np.cos(2*np.pi*z),axis=1))+20+np.e
        if self.name == "branin":
            a, b = 15*x[:, 0]-5, 15*x[:, 1]
            return (b-5.1*a*a/(4*np.pi**2)+5*a/np.pi-6)**2 + 10*(1-1/(8*np.pi))*np.cos(a)+10
        if self.name == "hartmann3":
            A = np.array([[3,10,30],[.1,10,35],[3,10,30],[.1,10,35]])
            P = 1e-4*np.array([[3689,1170,2673],[4699,4387,7470],[1091,8732,5547],[381,5743,8828]])
            return -np.sum(np.array([1,1.2,3,3.2])*np.exp(-np.sum(A*(x[:,None,:]-P)**2, axis=2)), axis=1)
        lo = np.array([.05,100,63070,990,63.1,700,1120,9855])
        hi = np.array([.15,50000,115600,1110,116,820,1680,12045])
        rw, r, tu, hu, tl, hl, length, kw = (lo+x*(hi-lo)).T
        lr = np.log(r/rw)
        return 2*np.pi*tu*(hu-hl)/(lr*(1+2*length*tu/(lr*rw**2*kw)+tu/tl))

    def base(self, x):

        if self.name in NEW_PROBLEMS:

            center, scale = {"sphere": (8*self.dim,8*self.dim),
                             "rosenbrock": (200*(self.dim-1),200*(self.dim-1)),
                             "ackley": (10.,10.)}[self.name[:-1]]
            return (self.raw(x)-center)/scale
        center, scale = {"branin": (50,50), "hartmann3": (-1,1),
                         "borehole": (80,50)}[self.name]
        return (self.raw(x)-center)/scale

    def delta(self, x, chain=1):
        x = np.atleast_2d(x)
        if self.stress:
            return self.discrepancy*(.7*np.sin(10*np.pi*x[:,0]) + .6*(x[:,1] > .5))
        if chain == 1:
            return self.discrepancy*(.18*np.sin(np.pi*x[:,0])+.12*(x[:,1]-.5)+.08*np.mean(x-.5, axis=1))
        return self.discrepancy*(-.12*np.cos(np.pi*x[:,1])+.10*x[:,0]*x[:,-1])

    def evaluate(self, x, level=2):
        if level not in (0, 1, 2):
            raise ValueError("level must be 0, 1, or 2")
        y = self.base(x)
        for chain in range(1, level+1):
            y = y+self.delta(x, chain)
        return y


import warnings
import numpy as np
from scipy.optimize import minimize
from sklearn.exceptions import ConvergenceWarning
from sklearn.gaussian_process import GaussianProcessRegressor
from sklearn.gaussian_process.kernels import ConstantKernel, Matern

def _p002(obj, initial, bounds):
    result = minimize(obj,initial,method="L-BFGS-B",jac=True,bounds=bounds,
                      options={"maxiter":45,"ftol":1e-7})
    return result.x, result.fun

class AdaptiveGP:

    def __init__(self):
        self.kernel = None
        self.last_optimized = -100

    def fit(self,x,y):
        x,y = np.asarray(x),np.asarray(y)
        kernel = self.kernel
        if kernel is None:
            kernel = ConstantKernel(1.,(.01,100))*Matern(np.full(x.shape[1],.35),(.04,4),nu=2.5)
        optimize = len(x)-self.last_optimized >= 8
        self.model = GaussianProcessRegressor(kernel=kernel,alpha=1e-6,normalize_y=True,
                                               optimizer=_p002 if optimize else None)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore",ConvergenceWarning)
            self.model.fit(x,y)
        self.kernel = self.model.kernel_
        if optimize:
            self.last_optimized = len(x)
        self.cache = {}
        return self

    def predict(self,x):
        x = np.atleast_2d(np.asarray(x,dtype=float))
        key = (x.shape,x.tobytes())
        if key not in self.cache:
            m,s = self.model.predict(x,return_std=True)
            if len(self.cache)>=5:
                self.cache.pop(next(iter(self.cache)))
            self.cache[key] = (m,np.maximum(s,1e-8))
        return self.cache[key]


import numpy as np
SEARCH_METHODS=("bapso","no_sobol","no_elite","boundary_pso","extreme_pso","sobol")

def _p003(task,level,budget,seed,method="bapso",population=20):
    if method not in SEARCH_METHODS:
        raise ValueError(method)
    rng=np.random.default_rng(seed)
    n=min(population,budget)
    use_sobol=method in ("bapso","no_elite","sobol")
    initial=_p001(n,task.dim,seed) if use_sobol else rng.random((n,task.dim))
    if method=="sobol":
        x=_p001(budget,task.dim,seed)
        return x,task.evaluate(x,level)
    x=initial; velocity=np.zeros_like(x)
    bx=x.copy();bf=np.full(n,np.inf);xs=[];ys=[]
    count=0
    while count<budget:
        m=min(n,budget-count);current=x[:m]
        y=task.evaluate(current,level);xs.append(current.copy());ys.append(y);count+=m
        f=y if method=="extreme_pso" else abs(y)
        ids=np.flatnonzero(f<bf[:m]);bx[ids]=current[ids];bf[ids]=f[ids]
        fraction=count/budget; best=bx[np.argmin(bf)]
        velocity=(.85-.45*fraction)*velocity+1.5*rng.random(x.shape)*(bx-x)
        if method in ("bapso","no_sobol"):
            k=max(2,int(n*(.5-.4*fraction)))
            elite=bx[np.argsort(bf)[:k]][rng.integers(k,size=n)]
            diversity=np.mean(np.linalg.norm(x-best,axis=1))/np.sqrt(task.dim)
            strength=min(1.,3*diversity)

            leader=(1-strength)*best+strength*elite
        else:
            leader=best
        velocity+=1.5*rng.random(x.shape)*(leader-x)
        x=np.clip(x+np.clip(velocity,-.2,.2),0,1)

        reset=rng.random(n)<.05
        x[reset]=rng.random((sum(reset),task.dim))
    return np.vstack(xs),np.concatenate(ys)


import numpy as np

class StudyTask(Problem):
    def __init__(self, name, family=0, stress=False):
        super().__init__(name,1.0,stress)
        self.family = family

    def delta(self, x, chain=1):
        x = np.atleast_2d(x)
        rng = np.random.default_rng(701+101*self.family+13*chain+self.dim)
        slope = rng.normal(size=self.dim)
        slope /= np.linalg.norm(slope)
        phase = rng.uniform(-np.pi,np.pi)
        smooth = .15*np.sin(np.pi*x[:,0]+phase)+.18*((x-.5)@slope)+.08*x[:,0]*x[:,-1]
        if self.stress:
            return smooth+.7*np.sin(10*np.pi*x[:,0]+phase)+.6*(x[:,1]>.5)
        return smooth


def _p004(model, x, truth, band=.15):
    mu, _ = model.predict(x)
    true_band, predicted_band = abs(truth) <= band, abs(mu) <= band
    true_class, predicted_class = truth <= 0, mu <= 0
    recalls = [np.mean(predicted_class[true_class] == True) if np.any(true_class) else np.nan,
               np.mean(predicted_class[~true_class] == False) if np.any(~true_class) else np.nan]
    return {"rmse": float(np.sqrt(np.mean((mu-truth)**2))),
            "balanced_accuracy": float(np.nanmean(recalls)),
            "boundary_recall": float(np.mean(predicted_band[true_band])) if np.any(true_band) else None,
            "boundary_precision": float(np.mean(true_band[predicted_band])) if np.any(predicted_band) else 0.,
            "boundary_mae": float(np.mean(abs(mu[true_band]-truth[true_band]))) if np.any(true_band) else None}


import argparse
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor,as_completed
from scipy.spatial.distance import cdist
from scipy.stats import norm
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parent
METHODS=('transfer','no_transfer')
NAMES={'transfer':'迁移测试','no_transfer':'非迁移测试'}

class IncrementalEpsilonSVR:


    def __init__(self,C=10.,epsilon=.005,gamma=1.,tol=1e-5,max_updates=200000):
        if C<=0 or epsilon<0 or gamma<=0 or tol<=0:raise ValueError('Invalid SVR parameters')
        self.C=float(C);self.epsilon=float(epsilon);self.gamma=float(gamma)
        self.tol=tol;self.max_updates=max_updates
        self.x=None;self.y=np.empty(0);self.beta=np.empty(0);self.bounds=np.empty(0)
        self.kernel=np.empty((0,0));self.gradient_=np.empty(0)
        self.intercept_=0.;self.pair_updates_=0;self.samples_seen_=0
        self.kernel_entries_computed_=0;self.kkt_gap_=0.
    def _kernel(self,x,z):return np.exp(-self.gamma*cdist(x,z,'sqeuclidean'))
    def _slopes(self):

        up=self.gradient_+self.epsilon*np.where(self.beta>=0,1.,-1.)
        down=self.gradient_+self.epsilon*np.where(self.beta>0,1.,-1.)
        can_up=self.beta<self.bounds-1e-10
        can_down=self.beta>-self.bounds+1e-10
        return np.where(can_up,up,np.inf),np.where(can_down,down,-np.inf)
    def _solve(self):
        if not np.any(self.bounds>0):
            self.intercept_=float(np.mean(self.y));self.kkt_gap_=0.
            self.support_=np.empty(0,dtype=int);self.support_vectors_=self.x[:0]
            self.dual_coef_=np.empty((1,0));return
        for _ in range(self.max_updates):
            up,down=self._slopes();i=int(np.argmin(up));j=int(np.argmax(down))
            self.kkt_gap_=max(0.,float(down[j]-up[i]))
            if self.kkt_gap_<=self.tol:break


            js=np.flatnonzero(np.isfinite(down));js=js[js!=i]
            bi=self.beta[i];bj=self.beta[js]
            cap=np.minimum(self.bounds[i]-bi,self.bounds[js]+bj)
            eta=np.maximum(self.kernel[i,i]+self.kernel[js,js]-2*self.kernel[i,js],0.)
            linear=self.gradient_[i]-self.gradient_[js]
            safe_eta=np.maximum(eta,1e-15)
            roots=np.array([-linear/safe_eta,
                            -(linear+2*self.epsilon)/safe_eta,
                            -(linear-2*self.epsilon)/safe_eta])
            candidates=np.vstack([np.zeros(len(js)),cap,
                                  np.full(len(js),-bi),bj,roots])
            candidates=np.clip(candidates,0.,cap)
            change=linear*candidates+.5*eta*candidates**2
            change+=self.epsilon*(abs(bi+candidates)+abs(bj-candidates)-abs(bi)-abs(bj))
            k,l=np.unravel_index(int(np.argmin(change)),change.shape)
            j=int(js[l]);step=float(candidates[k,l])
            if step<=0 or change[k,l]>=0:


                tiny=(abs(self.beta)<1e-12)&(self.beta!=0)
                if np.any(tiny):
                    self.beta[tiny]=0.
                    assert abs(self.beta.sum())<1e-8
                    self.gradient_=self.kernel@self.beta-self.y
                    continue
                raise RuntimeError('SVR numerical stagnation before KKT convergence')
            self.beta[i]+=step;self.beta[j]-=step
            self.gradient_+=step*(self.kernel[:,i]-self.kernel[:,j])
            self.pair_updates_+=1
        else:raise RuntimeError('Incremental SVR failed to meet KKT tolerance')
        up,down=self._slopes()
        lower=-float(np.min(up));upper=-float(np.max(down))
        if np.isfinite(lower) and np.isfinite(upper):self.intercept_=.5*(lower+upper)
        elif np.isfinite(lower):self.intercept_=lower
        elif np.isfinite(upper):self.intercept_=upper
        else:self.intercept_=float(np.mean(self.y))
        active=abs(self.beta)>1e-10
        self.support_=np.flatnonzero(active)
        self.support_vectors_=self.x[active]
        self.dual_coef_=self.beta[active][None,:]
    def partial_fit(self,x,y,sample_weight=None):
        x=np.atleast_2d(np.asarray(x,float));y=np.asarray(y,float).reshape(-1)
        weights=np.ones(len(y)) if sample_weight is None else np.asarray(sample_weight,float).reshape(-1)
        if len(x)!=len(y) or len(weights)!=len(y) or not np.isfinite(x).all() or not np.isfinite(y).all() or not np.isfinite(weights).all() or np.any(weights<0):
            raise ValueError('Invalid incremental observations or weights')
        if self.x is not None and x.shape[1]!=self.x.shape[1]:raise ValueError('Feature dimension changed')
        for point,response,weight in zip(x,y,weights):
            point=point[None,:];n=len(self.y)
            if n:
                cross=self._kernel(self.x,point).ravel()
                expanded=np.empty((n+1,n+1));expanded[:n,:n]=self.kernel
                expanded[:n,n]=cross;expanded[n,:n]=cross;expanded[n,n]=1.
                new_gradient=float(cross@self.beta-response)
                self.x=np.vstack([self.x,point]);self.kernel=expanded
            else:
                self.x=point.copy();self.kernel=np.ones((1,1));new_gradient=-float(response)
            self.y=np.append(self.y,response);self.beta=np.append(self.beta,0.)
            self.bounds=np.append(self.bounds,self.C*weight)
            self.gradient_=np.append(self.gradient_,new_gradient)
            self.samples_seen_+=1;self.kernel_entries_computed_+=n+1
            self._solve()
        return self
    def predict(self,x):
        if self.x is None:raise ValueError('SVR has not received observations')
        active=abs(self.beta)>1e-10
        return self._kernel(np.atleast_2d(x),self.x[active])@self.beta[active]+self.intercept_
    def gradient(self,x):
        x=np.atleast_2d(x);active=abs(self.beta)>1e-10
        delta=x[:,None,:]-self.x[active][None,:,:]
        kernel=np.exp(-self.gamma*np.sum(delta**2,axis=2))
        return np.einsum('ij,ijd,j->id',kernel,delta,-2*self.gamma*self.beta[active])
    def dual_objective(self):
        return float(.5*self.beta@self.kernel@self.beta+self.epsilon*abs(self.beta).sum()-self.y@self.beta)

class IncrementalSVR:


    def __init__(self,seed=0):self.seed=seed;self.x=None
    def _new(self):return IncrementalEpsilonSVR(C=10.,epsilon=.005,gamma=1./self.dim)
    def fit(self,x,y):
        x=np.asarray(x,float);y=np.asarray(y,float)
        if x.ndim!=2 or y.ndim!=1 or not len(y) or len(x)!=len(y):
            raise ValueError('Expected a nonempty observation matrix and response vector')
        if self.x is None:
            self.dim=x.shape[1];self.offset=float(np.mean(y));self.scale=max(float(np.std(y)),.05)
            self.model=self._new();self.members=[self._new() for _ in range(4)]
            self.rngs=[np.random.default_rng(self.seed+104729*(i+1)) for i in range(4)]
            start=0
        else:
            start=len(self.x)
            if len(x)<start or not np.array_equal(x[:start],self.x) or not np.array_equal(y[:start],self.y):
                raise ValueError('Incremental SVR accepts append-only data; old observations cannot change')
        for point,response in zip(x[start:],y[start:]):
            z=np.array([(response-self.offset)/self.scale]);point=point[None,:]
            self.model.partial_fit(point,z)
            for member,rng in zip(self.members,self.rngs):
                member.partial_fit(point,z,[rng.poisson(1.)])
        self.x=x.copy();self.y=y.copy()
        residual=y-(self.model.predict(self.x)*self.scale+self.offset)
        self.noise=max(float(np.sqrt(np.mean(residual**2))),.005*self.scale)
        return self
    def predict(self,x):
        x=np.atleast_2d(x)
        mean=self.model.predict(x)*self.scale+self.offset
        ensemble=np.array([m.predict(x)*self.scale+self.offset for m in self.members])
        distance=cdist(x,self.x).min(axis=1)/np.sqrt(self.dim)
        std=np.sqrt(ensemble.var(axis=0)+self.noise**2*(1+distance)**2)
        return mean,np.maximum(std,1e-8)
    def gradient(self,x):return self.scale*self.model.gradient(x)

class TransferSVR:

    def __init__(self,prior,seed=0):
        self.prior=prior;self.residual=IncrementalSVR(seed);self.target_y=np.empty(0)
    def fit(self,x,y):
        x=np.asarray(x,float);y=np.asarray(y,float);start=len(self.target_y)
        if len(y)<start or not np.array_equal(y[:start],self.target_y):
            raise ValueError('Historical target responses cannot change')


        old=np.empty(0) if not start else self.residual.y
        new=y[start:]-self.prior.predict(x[start:])[0] if len(y)>start else np.empty(0)
        self.residual.fit(x,np.concatenate([old,new]));self.target_y=y.copy();return self
    def predict(self,x):
        source,source_std=self.prior.predict(x);residual,residual_std=self.residual.predict(x)
        return source+residual,np.sqrt(source_std**2+residual_std**2)
    def gradient(self,x):return self.residual.gradient(x)

def _p005(values):
    values=np.asarray(values);return values/max(float(np.max(values)),1e-12)

def _p006(model,pool,observed,responses,prior_information=None,budget=40):


    distance=cdist(pool,observed).min(axis=1)
    mu,std=model.predict(pool)

    best=float(np.min(abs(responses)));improvement=best-abs(mu);z=improvement/std
    ei=improvement*norm.cdf(z)+std*norm.pdf(z)
    progress=len(observed)/budget

    band=abs(mu)<=.15
    coverage=float(np.mean(distance[band]/np.sqrt(pool.shape[1])<=.10)) if band.any() else 0.
    beta=.6*np.exp(-progress)*(1-coverage)
    gamma=.2*np.tanh(progress+coverage)
    score=_p005(ei)
    if prior_information is not None:
        prior_mu,prior_std=prior_information
        prior_boundary=np.exp(-.5*(prior_mu/np.maximum(prior_std,.15))**2)
        score+=beta*_p005(prior_boundary)
        gradient=model.gradient(pool)
    else:gradient=model.gradient(pool)
    score+=gamma*_p005(np.linalg.norm(gradient,axis=1))
    score*=np.minimum(distance/.08,1.)
    score[distance<=1e-8]=-np.inf
    if not np.isfinite(score).any():raise ValueError('Candidate pool exhausted')
    return pool[int(np.argmax(score))].copy()

def _p007(task,source,seed,method,prior,config,evaluation=None):
    x=_p001(8,task.dim,10000+seed);y=task.evaluate(x,source+1)
    pool=_p001(config['pool'],task.dim,20000+seed)
    if evaluation is None:
        test=_p001(config['test_size'],task.dim,100000+seed)
        truth=task.evaluate(test,source+1)
    else:test,truth=evaluation
    model=TransferSVR(prior,seed) if method=='transfer' else IncrementalSVR(seed)
    prior_information=prior.predict(pool) if method=='transfer' else None
    history=[]
    for n in range(8,config['target_budget']+1):
        model.fit(x,y)
        if (n-8)%4==0 or n==config['target_budget']:
            metrics=_p004(model,test,truth)
            if metrics['boundary_mae'] is None:raise ValueError('Empty independent boundary reference')
            history.append(dict(n=n,error=metrics['boundary_mae'],balanced_accuracy=metrics['balanced_accuracy']))
        if n==config['target_budget']:break
        point=_p006(model,pool,x,y,prior_information,config['target_budget'])
        x=np.vstack([x,point]);y=np.append(y,task.evaluate(point[None,:],source+1)[0])
    return dict(problem=task.name,source=source,seed=seed,method=method,x=x.tolist(),y=y.tolist(),history=history,
                target_queries=len(y),source_queries=config['source_budget'] if method=='transfer' else 0,
                signature=config['signature'])


def run_pair(job):
    problem,source,seed,config=job;task=StudyTask(problem,config['family'])
    sx,sy=_p003(task,source,config['source_budget'],seed,'bapso')
    prior=AdaptiveGP().fit(sx,sy)
    test=_p001(config['test_size'],task.dim,100000+seed);truth=task.evaluate(test,source+1)
    rows=[_p007(task,source,seed,method,prior,config,(test,truth)) for method in METHODS]
    for row in rows:
        x=np.asarray(row['x']);y=np.asarray(row['y'])
        assert len(y)==config['target_budget'] and len(np.unique(x,axis=0))==len(x)
        np.testing.assert_allclose(y,task.evaluate(x,source+1))
    np.testing.assert_array_equal(rows[0]['x'][:8],rows[1]['x'][:8])
    return rows


def _p008():
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,
        'axes.labelsize':9,'axes.titlesize':9,'xtick.labelsize':9,'ytick.labelsize':9,
        'legend.fontsize':9,'axes.edgecolor':'#333333','axes.linewidth':.8,
        'xtick.direction':'out','ytick.direction':'out','grid.color':'#D9D9D9',
        'grid.linewidth':.5,'grid.alpha':.6,'axes.axisbelow':True,
        'figure.facecolor':'white','axes.facecolor':'white','mathtext.fontset':'dejavusans'})

def plot_results(raw,config,path):
    _p008()
    problems=config['problems'];nrows=int(np.ceil(len(problems)/3))
    fig,axes=plt.subplots(nrows,3,figsize=(14.5,3.6*nrows),squeeze=False)
    for i,(ax,problem) in enumerate(zip(axes.ravel(),problems)):
        for method,label,color,linestyle in [('transfer','Transfer','#2B5C8A','-'),('no_transfer','No transfer','#B5443A','--')]:
            group=[r for r in raw if r['problem']==problem and r['method']==method]
            counts=[h['n'] for h in group[0]['history']]
            errors=np.array([[h['error'] for h in r['history']] for r in group])
            ax.plot(counts,errors.mean(axis=0),label=label,color=color,linestyle=linestyle,linewidth=1.3)
        ax.set_title(f'({chr(97+i)}) {problem}',loc='left',pad=7)
        ax.set(xlabel='Target-platform tests',ylabel='Boundary MAE',yscale='log');ax.grid()
    for ax in axes.ravel()[len(problems):]:ax.axis('off')
    handles,labels=axes.ravel()[0].get_legend_handles_labels()
    fig.legend(handles,labels,loc='lower center',ncol=2,frameon=False)
    fig.tight_layout(rect=(0,.055,1,1),pad=1.2,h_pad=1.5,w_pad=1.8)
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);fig.savefig(path,dpi=300);plt.close(fig)

def self_test():


    from sklearn.svm import SVR
    rng=np.random.default_rng(74);max_prediction=0.;max_objective=0.;prefixes=0
    for case in range(5):
        x=rng.random((32,3));y=np.sin(x[:,0]*4)+x[:,1]**2+.03*rng.normal(size=32)
        if case==1:y=np.full(32,2.3)
        if case==2:x[10:15]=x[0];y[10:15]=y[0]+np.linspace(-.3,.3,5)
        weights=rng.poisson(1.,32) if case==3 else np.ones(32)
        model=IncrementalEpsilonSVR(C=10,epsilon=.005,gamma=1/3)
        for n in range(1,33):
            old_kernel=model.kernel.copy()
            model.partial_fit(x[n-1:n],y[n-1:n],weights[n-1:n]);prefixes+=1
            np.testing.assert_array_equal(model.kernel[:n-1,:n-1],old_kernel)
            assert model.samples_seen_==n and model.kernel_entries_computed_==n*(n+1)//2
            assert abs(model.beta.sum())<1e-7 and np.all(abs(model.beta)<=10*weights[:n]+1e-7)
            assert model.kkt_gap_<=model.tol
            np.testing.assert_allclose(model.gradient_,model.kernel@model.beta-model.y,atol=1e-8)
            ids=np.flatnonzero(weights[:n]>0)
            if len(ids)<2:continue
            reference=SVR(C=10,epsilon=.005,gamma=1/3,tol=1e-8).fit(x[ids],y[ids],sample_weight=weights[ids])
            probe=rng.random((40,3))
            difference=float(np.max(abs(model.predict(probe)-reference.predict(probe))))
            beta=np.zeros(n);beta[ids[reference.support_]]=reference.dual_coef_[0]
            objective=.5*beta@model.kernel@beta+.005*abs(beta).sum()-y[:n]@beta
            max_prediction=max(max_prediction,difference)
            max_objective=max(max_objective,abs(model.dual_objective()-objective))
            assert difference<.002 and abs(model.dual_objective()-objective)<1e-5
    x=_p001(24,3,19);y=np.sin(3*x[:,0])+x[:,1]
    model=IncrementalSVR(23).fit(x[:8],y[:8]);scale=model.scale;offset=model.offset
    for n in range(9,25):model.fit(x[:n],y[:n])
    assert model.scale==scale and model.offset==offset
    assert all(m.samples_seen_==24 for m in [model.model]+model.members)
    before=model.model.pair_updates_;model.fit(x,y);assert model.model.pair_updates_==before
    try:model.fit(x,y+1)
    except ValueError:pass
    else:raise AssertionError('Changing old observations must be rejected')
    probe=_p001(8,3,24);gradient=model.gradient(probe);h=1e-5
    for j in range(3):
        offset=np.zeros(3);offset[j]=h
        finite=(model.predict(probe+offset)[0]-model.predict(probe-offset)[0])/(2*h)
        np.testing.assert_allclose(gradient[:,j],finite,rtol=1e-5,atol=1e-6)
    task=StudyTask('sphere6',3);sx,sy=_p003(task,0,24,3000,'bapso')
    prior=AdaptiveGP().fit(sx,sy);test=_p001(512,6,1234);truth=task.evaluate(test,1)
    class CountTask:
        name='sphere6';dim=6
        def __init__(self):self.count=0
        def evaluate(self,x,level):self.count+=len(x);return task.evaluate(x,level)
    config=dict(pool=64,target_budget=16,source_budget=24,signature='self-test')
    for method in METHODS:
        counted=CountTask();row=_p007(counted,0,3000,method,prior,config,(test,truth))
        assert counted.count==16 and len(np.unique(row['x'],axis=0))==16
        assert row==_p007(task,0,3000,method,prior,config,(test,truth))


    regression_x=np.array([[0.9775526691228151, 0.18838719371706247], [0.34642961248755455, 0.7689927769824862], [0.1741474810987711, 0.35033193603157997], [0.5565732605755329, 0.6769072916358709], [0.7482740050181746, 0.4469564203172922], [0.11477056983858347, 0.525591678917408], [0.41154356207698584, 0.10935606714338064], [0.7957393536344171, 0.9339649127796292]])
    regression_y=np.array([-1.031318005743752, 0.25247991760884003, -0.2839811112911058, 0.4863718670332795, 0.12824927143812803, -0.4510558473896284, -0.4598342751167065, 2.800526717988111])
    regression=IncrementalSVR(3015).fit(regression_x,regression_y)
    assert all(member.kkt_gap_<=member.tol for member in [regression.model]+regression.members)
    print(f'PASS: {prefixes} prefixes; max batch prediction error={max_prediction:.9g}; max dual-objective difference={max_objective:.9g}')
    print('PASS: KKT, box/equality constraints, retained kernel/state, frozen normalization, gradients, exact queries and deterministic replay.')


def _write_book(path,tables):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    import math,json
    workbook=Workbook(write_only=True)
    header_font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
    body_font=Font(name='Arial',size=10,color='222222')
    fill=PatternFill('solid',fgColor='2B5C8A')
    border=Border(bottom=Side(style='thin',color='CCCCCC'))
    header_alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    number_alignment=Alignment(horizontal='right',vertical='center')
    text_alignment=Alignment(horizontal='left',vertical='center')
    wrap_alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
    for title,records in tables:
        if not records:raise ValueError('Empty export table: '+title)
        sheet=workbook.create_sheet(title);headers=list(records[0])
        sheet.sheet_view.showGridLines=False
        sheet.freeze_panes='C2' if title in ('Trials','Curves','Observations') else 'A2'
        sheet.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{len(records)+1}'
        sheet.row_dimensions[1].height=42
        for column,key in enumerate(headers,1):
            sample=[r.get(key) for r in records[:100]]
            width=max(14,min(34,max([len(key.replace('_',' '))*.8]+[len(str(v)) for v in sample if isinstance(v,str)])))
            sheet.column_dimensions[get_column_letter(column)].width=90 if key=='value' else width
        cells=[]
        for value in headers:
            cell=WriteOnlyCell(sheet,value=value);cell.font=header_font;cell.fill=fill
            cell.alignment=header_alignment;cell.border=border;cells.append(cell)
        sheet.append(cells)
        for index,record in enumerate(records,2):
            cells=[]
            for key in headers:
                value=record.get(key)
                if isinstance(value,np.generic):value=value.item()
                if isinstance(value,(list,tuple,dict)):value=json.dumps(value,ensure_ascii=True)
                if isinstance(value,float) and not math.isfinite(value):raise ValueError('Non-finite export value')
                if isinstance(value,str) and not value.isascii():raise ValueError('Workbook text must be English ASCII')
                cell=WriteOnlyCell(sheet,value=value);cell.font=body_font
                number=isinstance(value,(int,float)) and not isinstance(value,bool)
                cell.alignment=number_alignment if number else text_alignment
                if key=='value' and isinstance(value,str) and len(value)>80:
                    cell.alignment=wrap_alignment;sheet.row_dimensions[index].height=15*int(np.ceil(len(value)/80))
                if isinstance(value,float):cell.number_format='0.00%' if 'coverage' in key or 'accuracy' in key else '0.00000'
                elif isinstance(value,int):cell.number_format='0'
                cells.append(cell)
            sheet.append(cells)
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize='9'
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:1'
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);workbook.save(path)

def _settings(values):
    import hashlib
    return [dict(parameter='source_script',value=Path(__file__).name),
            dict(parameter='source_sha256',value=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())]+[
            dict(parameter=key,value=value) for key,value in values.items()]


def export_results(raw,config):
    summary=[];curves=[];observations=[]
    for problem in config['problems']:
        for method in METHODS:
            group=[r for r in raw if r['problem']==problem and r['method']==method]
            for n in [h['n'] for h in group[0]['history']]:
                values=np.array([next(h['error'] for h in r['history'] if h['n']==n) for r in group])
                accuracy=np.array([next(h['balanced_accuracy'] for h in r['history'] if h['n']==n) for r in group])
                summary.append(dict(problem=problem,method=method,target_tests=n,repeats=len(group),
                    boundary_mae_mean=float(values.mean()),boundary_mae_sd=float(values.std(ddof=1)),
                    balanced_accuracy_mean=float(accuracy.mean())))
    for r in raw:
        for h in r['history']:
            curves.append(dict(problem=r['problem'],source_level=r['source'],seed=r['seed'],method=r['method'],
                target_tests=h['n'],boundary_mae=h['error'],balanced_accuracy=h['balanced_accuracy']))
        for index,(x,y) in enumerate(zip(r['x'],r['y']),1):
            record=dict(problem=r['problem'],source_level=r['source'],seed=r['seed'],method=r['method'],
                        query=index,response=y,dimension=len(x))
            record.update({f'x{j+1}':float(x[j]) if j<len(x) else None for j in range(8)})
            observations.append(record)
    settings=_settings({**config,'methods':list(METHODS),
        'target_model':'Incremental RBF epsilon-SVR','transfer_model':'Source GP plus incremental SVR residual',
        'uncertainty':'Four persistent Poisson bootstrap models','metric_direction':'Lower boundary MAE is better'})
    _write_book(ROOT/'output'/'03_output.xlsx',[
        ('Summary',summary),('Curves',curves),('Observations',observations),('Settings',settings)])

def main():
    parser=argparse.ArgumentParser(description='Transfer vs no transfer; saves PNG and XLSX results')
    parser.add_argument('--self-test',action='store_true',help='Check incremental solver against batch SVR; save no files')
    parser.add_argument('--seeds',type=int,default=20);parser.add_argument('--seed-start',type=int,default=3000)
    parser.add_argument('--family',type=int,default=3);parser.add_argument('--budget',type=int,default=40)
    parser.add_argument('--low-budget',type=int,default=160);parser.add_argument('--pool',type=int,default=512)
    parser.add_argument('--workers',type=int,default=8);parser.add_argument('--problems',nargs='+',default=PROBLEMS,choices=PROBLEMS)
    args=parser.parse_args()
    if args.self_test:return self_test()
    if args.seeds<2 or args.budget<16 or args.pool<args.budget or args.low_budget<20 or args.workers<1:parser.error('seeds>=2, budget>=16, pool>=budget, source budget>=20, workers>=1 required')
    config=dict(problems=args.problems,seeds=list(range(args.seed_start,args.seed_start+args.seeds)),family=args.family,
                target_budget=args.budget,source_budget=args.low_budget,pool=args.pool,test_size=4096,signature='in-memory-public-reference')
    jobs=[(problem,source,seed,config) for problem in args.problems for seed in config['seeds'] for source in (0,1)]
    raw=[]
    if args.workers==1:
        for job in jobs:raw.extend(run_pair(job));print(job[:3],'complete',flush=True)
    else:
        with ProcessPoolExecutor(max_workers=args.workers) as executor:
            for i,result in enumerate(executor.map(run_pair,jobs),1):raw.extend(result);print(f'Completed {i}/{len(jobs)}',flush=True)
    plot_results(raw,config,ROOT/'output'/'03_output.png')
    for problem in config['problems']:
        for n in (16,config['target_budget']):
            means={m:float(np.mean([next(h['error'] for h in r['history'] if h['n']==n)
                for r in raw if r['problem']==problem and r['method']==m])) for m in METHODS}
            print(f"RESULT {problem} n={n} transfer={means['transfer']:.8f} no_transfer={means['no_transfer']:.8f}",flush=True)
    export_results(raw,config)
    print('Saved PNG and XLSX:',ROOT/'output')

if __name__=='__main__':main()
