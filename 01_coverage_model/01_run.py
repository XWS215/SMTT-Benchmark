


from __future__ import annotations
import os
import colorsys
import sys
from pathlib import Path
import numpy as np
from scipy.interpolate import make_smoothing_spline, UnivariateSpline
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.colors import to_rgb
import pandas as pd
from scipy.stats import spearmanr


EPS = 1e-12


def _p001(c_den, c_uni, c_acc, c_geo,
                  w_den=0.1, w_uni=0.3, w_acc=0.5, w_geo=-0.1):

    vals = [c_den, c_uni, c_acc, c_geo]
    ws = [w_den, w_uni, w_acc, w_geo]
    with np.errstate(divide='ignore'):
        logv = np.log(np.clip(vals, EPS, None))
    return float(np.exp(ws[0] * logv[0] + ws[1] * logv[1] +
                        ws[2] * logv[2] + ws[3] * logv[3]))


def _p002(c_den, c_uni, c_acc, c_geo,
                     w_den=0.1, w_uni=0.3, w_acc=0.5, w_geo=-0.1):

    return float(w_den * c_den + w_uni * c_uni + w_acc * c_acc + w_geo * c_geo)


def _p003(x_eval, y_eval):

    t = np.asarray(x_eval, float)
    y = np.asarray(y_eval, float)
    n = len(t)
    if n < 3:
        return 1.0
    dev = []
    for i in range(n - 2):
        v1 = np.array([t[i + 1] - t[i], y[i + 1] - y[i]])
        v2 = np.array([t[i + 2] - t[i + 1], y[i + 2] - y[i + 1]])
        n1 = np.linalg.norm(v1) + EPS
        n2 = np.linalg.norm(v2) + EPS
        ang = np.arccos(np.clip(np.dot(v1 / n1, v2 / n2), -1.0, 1.0))
        dev.append(ang)
    return float(np.mean(dev))


def _p004(x_sample, eps=EPS):

    x = np.sort(np.asarray(x_sample, float))
    n = len(x)
    if n < 3:
        return 0.0
    dx = np.diff(x)
    if dx.mean() <= 0:
        return 0.0
    cv = dx.std() / (dx.mean() + eps)
    return float(np.exp(-cv))


def _p005(points_2d, eps=EPS):


    pts = np.asarray(points_2d, float)
    n = len(pts)
    if n < 3:
        return 0.0
    o = np.argsort(pts[:, 0])
    xs, ys = pts[o, 0], pts[o, 1]
    width = np.zeros(n)
    width[0] = (xs[1] - xs[0])
    width[1:-1] = (xs[2:] - xs[:-2]) / 2.0
    width[-1] = (xs[-1] - xs[-2])
    yspan = ys.max() - ys.min()
    if yspan < EPS:
        yspan = 1.0
    area = np.maximum(width, EPS) * yspan
    atot = area.sum()
    if atot <= 0:
        return 0.0
    return float(1.0 - area.max() / atot)


def _p006(true_label, pred_label):

    true_label = np.asarray(true_label)
    pred_label = np.asarray(pred_label)
    if len(true_label) == 0:
        return 0.5
    return float(np.mean(pred_label == true_label))


def _p007(x_bound, y_bound, spl, x_val, y_val, true_label,
                        uni_mode='spacing', n_eval_pts=200):


    x_b = np.asarray(x_bound, float)
    y_b = np.asarray(y_bound, float)
    n_s = len(x_b)

    c_den = n_s / 100.0

    if spl is not None and n_s >= 3:
        t_eval = np.linspace(x_b.min(), x_b.max(), n_eval_pts)
        y_eval = spl(t_eval)
        c_geo = _p003(t_eval, y_eval)
    else:
        c_geo = 1.0

    if uni_mode == 'voronoi':
        c_uni = _p005(np.column_stack([x_b, y_b]))
    else:
        c_uni = _p004(x_b)

    if spl is not None:
        y_val_fit = spl(np.asarray(x_val, float))
        pred = (np.asarray(y_val, float) > y_val_fit).astype(int)
        c_acc = _p006(true_label, pred)
    else:
        c_acc = 0.5
    return dict(c_geo=c_geo, c_den=c_den, c_uni=c_uni, c_acc=c_acc)


def _p008(m, form='exp', **w):

    if form == 'linear':
        return _p002(m['c_den'], m['c_uni'], m['c_acc'], m['c_geo'], **w)
    return _p001(m['c_den'], m['c_uni'], m['c_acc'], m['c_geo'], **w)


def _p009(x, y, p=0.6, min_pts=4):


    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    if len(x) < min_pts or len(x) != len(y):
        return None
    order = np.argsort(x)
    x, y = x[order], y[order]


    ux, inv = np.unique(x, return_inverse=True)
    if len(ux) < len(x):
        ysum = np.bincount(inv, weights=y)
        ycnt = np.bincount(inv)
        x = ux
        y = ysum / np.maximum(ycnt, 1)
    if len(x) < min_pts:
        return None
    p = float(np.clip(p, 1e-6, 1.0 - 1e-6))
    lam = (1.0 - p) / p
    try:
        spl = make_smoothing_spline(x, y, lam=lam)
        return spl
    except (ValueError, TypeError):

        s = min(len(x) * 0.05, (y.max() - y.min()) ** 2)
        return UnivariateSpline(x, y, k=3, s=s)


def _p010(x, y, xx, p=0.6):

    spl = _p009(x, y, p=p)
    xx = np.asarray(xx, dtype=float)
    if spl is None:
        return np.full_like(xx, np.nan, dtype=float)
    return spl(xx)


CM2IN = 1.0 / 2.54


COLOR_A = np.array([0.000, 0.447, 0.741])       
COLOR_B = np.array([0.850, 0.325, 0.098])       
COLOR_C = np.array([0.466, 0.674, 0.188])       
COLOR_D = np.array([0.494, 0.184, 0.556])       
COLOR_E = np.array([0.635, 0.078, 0.184])        
COLOR_F = np.array([0.302, 0.745, 0.933])        
COLOR_G = np.array([0.929, 0.694, 0.125])        
_GROUP_COLORS = dict(A=COLOR_A, B=COLOR_B, C=COLOR_C, D=COLOR_D,
                     E=COLOR_E, F=COLOR_F, G=COLOR_G)


def _p011(name):

    return _GROUP_COLORS.get(name, COLOR_A)


ACADEMIC_HEX = {
    "A": "#5878A6",             
    "B": "#F4A7AF",        
    "C": "#A6D6C6",         
    "D": "#A7C3E7",         
    "E": "#FBD9C9",        
    "F": "#CFEAF6",         
    "G": "#DADDE3",        
}


def _p012(name, role="line"):


    c = to_rgb(ACADEMIC_HEX.get(name, ACADEMIC_HEX["A"]))
    h, s, v = colorsys.rgb_to_hsv(*c)
    if role == "line":
        s, v = min(1.0, s * 1.35), v * 0.78
    else:
        s, v = min(1.0, s * 1.15), v * 0.90
    return np.array(colorsys.hsv_to_rgb(h, s, v))


COLORS_POP = [COLOR_A, COLOR_B, COLOR_C, COLOR_D]
COLORS_XFER = [COLOR_A, COLOR_B, COLOR_C, COLOR_D,
               COLOR_E, COLOR_F, COLOR_G]


def _p013():
    try:
        candidates = ["Times New Roman", "Times", "Liberation Serif", "DejaVu Serif"]
        for name in candidates:
            if any(f.name == name for f in font_manager.fontManager.ttflist):
                plt.rcParams["font.family"] = name
                return name
    except Exception:
        pass
    plt.rcParams["font.family"] = "serif"
    return "serif"


def _p014():
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,
        'axes.labelsize':9,'axes.titlesize':9,'xtick.labelsize':9,'ytick.labelsize':9,
        'legend.fontsize':9,'axes.edgecolor':'#333333','axes.linewidth':.8,
        'xtick.direction':'out','ytick.direction':'out','grid.color':'#D9D9D9',
        'grid.linewidth':.5,'grid.alpha':.6,'axes.axisbelow':True,
        'figure.facecolor':'white','axes.facecolor':'white','mathtext.fontset':'dejavusans'})

_p014()


def _p015(ax, fs=9, grid=True, grid_alpha=0.6):
    ax.set_facecolor("white")
    for s in ax.spines.values():
        s.set_color("#333333")
        s.set_linewidth(0.8)
    ax.tick_params(colors="k")
    for lbl in ax.get_xticklabels() + ax.get_yticklabels():
        lbl.set_fontsize(fs)
    if grid:
        ax.grid(True, alpha=grid_alpha, linestyle="-", linewidth=0.5)
        ax.set_axisbelow(True)
    try:                                                             
        ax.ticklabel_format(useOffset=False)
    except Exception:
        pass


def _p016(w_cm=10, h_cm=5):
    fig = plt.figure(figsize=(w_cm * CM2IN, h_cm * CM2IN))
    return fig


def _p017(fig, path, dpi=600, close=True, extra=("png",), tight=True):


    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    base = os.path.splitext(path)[0]
    fmt = os.path.splitext(path)[1].lstrip(".") or "png"
    fmts = list(dict.fromkeys([fmt] + list(extra)))
    bbox = "tight" if tight else None
    for f in fmts:
        try:
            fig.savefig(f"{base}.{f}", dpi=dpi, bbox_inches=bbox,
                        facecolor="white", backend=None)
            print(f"  saved {base}.{f}")
        except Exception as e:
            print(f"  !! save {f} failed: {e}")
    if close:
        plt.close(fig)


def _p018(ax, x, mean, std, color, alpha=0.22, label=None):
    m = np.asarray(mean)
    s = np.asarray(std)
    ax.plot(x, m, color=color, lw=1.4, label=label)
    ax.fill_between(np.asarray(x), m - s, m + s, color=color, alpha=alpha,
                    linewidth=0)


def _p019(x):
    return 3 + 4 * np.sin(0.5 * x) + 0.3 * x + 0.5 * np.cos(1.2 * x)


def _p020(seed=2691, n1=200, n2=50, n_val=2000):

    rng = np.random.default_rng(seed)
    x_val = 10.0 * rng.random(n_val)
    y_val = 10.0 * rng.random(n_val)
    true_label = (y_val > _p019(x_val)).astype(int)
    x1 = 10.0 * rng.random(n1)
    y1 = 10.0 * rng.random(n1)
    x2 = 10.0 * rng.random(n2)
    y2 = 10.0 * rng.random(n2)
    return dict(x_val=x_val, y_val=y_val, true_label=true_label,
                x1=x1, y1=y1, x2=x2, y2=y2)


def _p021(xs, ys, band):
    sel = np.abs(ys - _p019(xs)) < band
    return np.asarray(xs[sel], float), np.asarray(ys[sel], float)


def _p022(ds, band, p=0.6, uni_mode='spacing'):

    outs = []
    for xs, ys in ((ds['x1'], ds['y1']), (ds['x2'], ds['y2'])):
        xb, yb = _p021(xs, ys, band)
        spl = _p009(xb, yb, p=p)
        m = _p007(xb, yb, spl, ds['x_val'], ds['y_val'],
                                ds['true_label'], uni_mode=uni_mode)
        m['n_bound'] = len(xb)
        outs.append(m)
    return outs[0], outs[1]


def _p023(m, **w):
    return _p001(m['c_den'], m['c_uni'], m['c_acc'], m['c_geo'], **w)


def _p024(m, **w):
    return _p002(m['c_den'], m['c_uni'], m['c_acc'], m['c_geo'], **w)


SAMPLE_STRATEGIES = (
    dict(key="S1", kind="uniform", n_exec=400, tier="sufficient",
         en="S1 Uniform (400 exec)", short="S1 uniform, 400",
         zh="均匀密采（400 执行）"),
    dict(key="S2", kind="uniform", n_exec=200, tier="sufficient",
         en="S2 Uniform (200 exec)", short="S2 uniform, 200",
         zh="均匀中采（200 执行）"),
    dict(key="S3", kind="uniform", n_exec=100, tier="critical",
         en="S3 Uniform (100 exec)", short="S3 uniform, 100",
         zh="均匀疏采（100 执行）"),
    dict(key="S4", kind="cluster", n_exec=200, tier="degraded",
         en="S4 Two-cluster (200 exec)", short="S4 two-cluster, 200",
         zh="双簇聚集（200 执行）"),
    dict(key="S5", kind="gap", n_exec=200, tier="degraded",
         en="S5 Mid-gap (200 exec)", short="S5 mid-gap, 200",
         zh="中段空洞（200 执行）"),
)


def _p025(rng, n_val=2000):

    x_val = 10.0 * rng.random(n_val)
    y_val = 10.0 * rng.random(n_val)
    return dict(x_val=x_val, y_val=y_val,
                true_label=(y_val > _p019(x_val)).astype(int))


def _p026(rng, kind="uniform", n_exec=200, band=1.5):

    n = int(n_exec)
    if kind == "uniform":
        x = 10.0 * rng.random(n)
    elif kind == "cluster":
        n_l = n // 2
        x = np.r_[1.5 + 2.0 * rng.random(n_l),
                  6.5 + 2.0 * rng.random(n - n_l)]
    elif kind == "gap":
        chunks, have = [], 0
        while have < n:
            xc = 10.0 * rng.random(2 * n)
            keep = xc[(xc < 3.5) | (xc > 6.5)]
            chunks.append(keep)
            have += len(keep)
        x = np.concatenate(chunks)[:n]
    else:
        raise ValueError(f"unknown strategy kind: {kind}")
    y = 10.0 * rng.random(n)
    xb, yb = _p021(x, y, band)
    return x, y, xb, yb


def _p027(xb, yb, ds, p=0.6, uni_mode='spacing'):

    spl = _p009(xb, yb, p=p)
    m = _p007(xb, yb, spl, ds['x_val'], ds['y_val'],
                            ds['true_label'], uni_mode=uni_mode)
    m['n_bound'] = len(xb)
    return m


def _p028(x_bound, eps=0.5, m=3000):


    if len(x_bound) == 0:
        return 0.0
    xt = np.linspace(0.0, 10.0, m)
    yt = _p019(xt)
    xs = np.asarray(x_bound, float)
    d = np.min(np.hypot(xt[:, None] - xs[None, :],
                        yt[:, None] - _p019(xs)[None, :]), axis=1)
    return float(np.mean(d <= eps))


ROOT = Path(__file__).resolve().parent
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


OUT_FIG = ROOT / "output"


BAND = 1.5                                   
SEED_MAIN = 2691            
SEEDS = [2691, 2692, 2693, 2694, 2695]            
TAU = 0.95                             
EPS_TRUE = 0.5                 
METRIC_KEYS = ["c_den", "c_uni", "c_acc", "c_geo"]
METRIC_LABEL = {"c_den": r"$C_{den}$", "c_uni": r"$C_{uni}$",
                "c_acc": r"$C_{acc}$", "c_geo": r"$C_{geo}$"}

STYLE = {
    "S1": dict(color="#2B5C8A", marker="o"),
    "S2": dict(color="#5B8FC9", marker="s"),
    "S3": dict(color="#9CB9D0", marker="^"),
    "S4": dict(color="#B5443A", marker="X"),
    "S5": dict(color="#D4A15A", marker="D"),
}


def run_seed(seed):


    rng = np.random.default_rng(seed)
    ds = _p025(rng)
    out = []
    for st in SAMPLE_STRATEGIES:
        _x, _y, xb, yb = _p026(rng, st["kind"], st["n_exec"], BAND)
        m = _p027(xb, yb, ds)
        m_v = _p027(xb, yb, ds, uni_mode="voronoi")
        out.append((st, m, _p028(xb, eps=EPS_TRUE), len(xb),
                    _p023(m_v), m_v["c_uni"]))
    return out


def _write_book(path,tables):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    import math,json
    workbook=Workbook(write_only=True)
    header_font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
    body_font=Font(name='Arial',size=10,color='222222')
    fill=PatternFill('solid',fgColor='2B5C8A')
    border=Border(bottom=Side(style='thin',color='CCCCCC'))
    header_alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    number_alignment=Alignment(horizontal='right',vertical='center')
    text_alignment=Alignment(horizontal='left',vertical='center')
    wrap_alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
    for title,records in tables:
        if not records:raise ValueError('Empty export table: '+title)
        sheet=workbook.create_sheet(title);headers=list(records[0])
        sheet.sheet_view.showGridLines=False
        sheet.freeze_panes='C2' if title in ('Trials','Curves','Observations') else 'A2'
        sheet.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{len(records)+1}'
        sheet.row_dimensions[1].height=42
        for column,key in enumerate(headers,1):
            sample=[r.get(key) for r in records[:100]]
            width=max(14,min(34,max([len(key.replace('_',' '))*.8]+[len(str(v)) for v in sample if isinstance(v,str)])))
            sheet.column_dimensions[get_column_letter(column)].width=90 if key=='value' else width
        cells=[]
        for value in headers:
            cell=WriteOnlyCell(sheet,value=value);cell.font=header_font;cell.fill=fill
            cell.alignment=header_alignment;cell.border=border;cells.append(cell)
        sheet.append(cells)
        for index,record in enumerate(records,2):
            cells=[]
            for key in headers:
                value=record.get(key)
                if isinstance(value,np.generic):value=value.item()
                if isinstance(value,(list,tuple,dict)):value=json.dumps(value,ensure_ascii=True)
                if isinstance(value,float) and not math.isfinite(value):raise ValueError('Non-finite export value')
                if isinstance(value,str) and not value.isascii():raise ValueError('Workbook text must be English ASCII')
                cell=WriteOnlyCell(sheet,value=value);cell.font=body_font
                number=isinstance(value,(int,float)) and not isinstance(value,bool)
                cell.alignment=number_alignment if number else text_alignment
                if key=='value' and isinstance(value,str) and len(value)>80:
                    cell.alignment=wrap_alignment;sheet.row_dimensions[index].height=15*int(np.ceil(len(value)/80))
                if isinstance(value,float):cell.number_format='0.00%' if 'coverage' in key or 'accuracy' in key else '0.00000'
                elif isinstance(value,int):cell.number_format='0'
                cells.append(cell)
            sheet.append(cells)
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize='9'
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:1'
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);workbook.save(path)

def _settings(values):
    import hashlib
    return [dict(parameter='source_script',value=Path(__file__).name),
            dict(parameter='source_sha256',value=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())]+[
            dict(parameter=key,value=value) for key,value in values.items()]


def export_results(df,rho_pool):
    trials=df.drop(columns=['zh']).rename(columns={'en':'strategy_label'}).to_dict('records')
    summary=[]
    for strategy in [r['key'] for r in SAMPLE_STRATEGIES]:
        group=df[df['strategy']==strategy]
        record=dict(strategy=strategy,repeats=len(group),executions=int(group['n_exec'].iloc[0]))
        for key in ('n_bound','c_den','c_uni','c_acc','c_geo','c_cover','truth_cov'):
            record[key+'_mean']=float(group[key].mean());record[key+'_sd']=float(group[key].std(ddof=1))
        summary.append(record)
    settings=_settings(dict(seeds=SEEDS,plot_seed=SEED_MAIN,boundary_band=BAND,score_threshold=TAU,
        reference_radius=EPS_TRUE,pooled_spearman=rho_pool,
        reference_definition='Equal-x-grid coverage after projection onto the known curve'))
    _write_book(OUT_FIG/'01_output.xlsx',[('Summary',summary),('Trials',trials),('Settings',settings)])

def main():
    for d in (OUT_FIG,):
        d.mkdir(parents=True, exist_ok=True)


    rows = []
    for s in SEEDS:
        for st, m, tcov, nb, ccov_vor, cuni_vor in run_seed(s):
            rows.append(dict(strategy=st["key"], en=st["en"], zh=st["zh"],
                             tier=st["tier"], n_exec=st["n_exec"],
                             seed=s, n_bound=nb,
                             c_geo=m["c_geo"], c_den=m["c_den"],
                             c_uni=m["c_uni"], c_acc=m["c_acc"],
                             c_cover=_p023(m), truth_cov=tcov,
                             c_uni_vor=cuni_vor, c_cover_vor=ccov_vor))
    df = pd.DataFrame(rows)

    keys = [st["key"] for st in SAMPLE_STRATEGIES]
    agg = (df.groupby("strategy")[["n_exec", "n_bound"] + METRIC_KEYS
                                  + ["c_cover", "truth_cov"]]
             .agg(["mean", "std"]).reindex(keys))


    rho = {}
    for s in SEEDS:
        d = df[df["seed"] == s]
        rho[s] = float(spearmanr(d["c_cover"], d["truth_cov"]).correlation)
    rho_pool = float(spearmanr(df["c_cover"], df["truth_cov"]).correlation)
    rho_vor = {s: float(spearmanr(df[df["seed"] == s]["c_cover_vor"],
                                  df[df["seed"] == s]["truth_cov"]).correlation)
               for s in SEEDS}


    deg = df[df["tier"] == "degraded"]["c_cover"]
    suf = df[df["tier"] == "sufficient"]["c_cover"]
    cri = df[df["tier"] == "critical"]["c_cover"]
    sep_lo, sep_hi = float(deg.max()), float(suf.min())
    degraded_below_tau = int((deg < TAU).sum())
    sufficient_above_tau = int((suf >= TAU).sum())
    critical_above_tau = int((cri >= TAU).sum())


    print("\n=========== RQ1-1  B-GCAM validation on synthetic boundary ===========")
    print(f"band = {BAND}, 5 sampling strategies x {len(SEEDS)} seeds")
    hdr = (f"{'Strategy':<26s}{'N_b':>12s} " + " ".join(
        f"{METRIC_LABEL[k]:>16s}" for k in METRIC_KEYS)
        + f" {'C_cover':>17s}{'true cov':>17s}")
    print(hdr)
    for k in keys:
        r = agg.loc[k]
        cells = " ".join(f"{r[(c, 'mean')]:6.3f}±{r[(c, 'std')]:.3f}"
                         for c in METRIC_KEYS)
        print(f"{k:<26s} {r[('n_bound', 'mean')]:5.1f}±{r[('n_bound', 'std')]:4.1f} "
              f"{cells} {r[('c_cover', 'mean')]:6.3f}±{r[('c_cover', 'std')]:.3f} "
              f"{r[('truth_cov', 'mean')]:6.3f}±{r[('truth_cov', 'std')]:.3f}")
    print(f"\nSpearman(C_cover, true coverage): "
          + ", ".join(f"{s}:{v:.2f}" for s, v in rho.items())
          + f";  pooled(25 pts)={rho_pool:.3f}")
    print(f"gap: degraded max={sep_lo:.3f} < sufficient min={sep_hi:.3f}; "
          f"tau={TAU} in gap -> degraded below tau {degraded_below_tau}/{len(deg)}, "
          f"sufficient above tau {sufficient_above_tau}/{len(suf)}, "
          f"critical above tau {critical_above_tau}/{len(cri)}")
    print("Voronoi-variant Spearman: "
          + ", ".join(f"{s}:{v:.2f}" for s, v in rho_vor.items())
          + "  (paper eq.(27)-(29) definition; does NOT separate S4)")


    FS = 9.0                                   
    fig = _p016(25.0, 16.5)
    x_fit = np.linspace(0, 10, 401)
    y_true = _p019(x_fit)


    rng = np.random.default_rng(SEED_MAIN)
    ds_main = _p025(rng)
    panels = []
    for st in SAMPLE_STRATEGIES:
        x, y, xb, yb = _p026(rng, st["kind"], st["n_exec"], BAND)
        panels.append((st, x, y, xb, yb, _p027(xb, yb, ds_main)))

    for k, (st, x, y, xb, yb, m) in enumerate(panels):
        ax = fig.add_subplot(2, 3, k + 1)
        inband = np.abs(y - _p019(x)) < BAND
        ax.scatter(x[~inband], y[~inband], s=6, color="#cfcfcf", marker="o",
                   linewidths=0, label="Non-boundary")
        ax.scatter(xb, yb, s=20, color="#2B5C8A", marker="o",
                   edgecolors="#333333", linewidths=0.35, label="Boundary")
        ax.plot(x_fit, y_true, "k-", lw=1.3, label="True boundary")
        spl = _p009(xb, yb, p=0.6)
        if spl is not None:
            ax.plot(x_fit, spl(x_fit), "--", color="#B5443A", lw=1.3, label="Fitted boundary")
        ax.set_xlim(0, 10)
        ax.set_ylim(0, 10)
        ax.set_xlabel(r"Parameter $x_1$", fontsize=FS)
        ax.set_ylabel(r"Parameter $x_2$", fontsize=FS)
        tcov = _p028(xb, eps=EPS_TRUE)
        ax.set_title(f"({chr(97 + k)}) {st['en']}", loc='left', fontsize=FS)
        _p015(ax, fs=FS)
        if k == 3:                              
            ax.legend(fontsize=FS, loc="lower left", frameon=True,
                      framealpha=0.92, handletextpad=0.4, labelspacing=0.35,
                      borderpad=0.35)


    ax = fig.add_subplot(2, 3, 6)
    for st in SAMPLE_STRATEGIES:
        d = df[df["strategy"] == st["key"]]
        sty = STYLE[st["key"]]
        ax.scatter(d["c_cover"], d["truth_cov"], s=34, color=sty["color"],
                   marker=sty["marker"], edgecolors="#333333", linewidths=0.4,
                   label=st["key"], zorder=3)
    ax.axvline(TAU, color="#666666", lw=1.0, ls="--", zorder=1)
    ax.text(TAU - 0.010, 0.437, rf"$\tau = {TAU}$", ha="right", va="bottom",
            fontsize=FS, color="#444444")
    ax.set_xlabel(r"$C_{cover}$", fontsize=FS)
    ax.set_ylabel("True boundary coverage", fontsize=FS)
    ax.set_title('(f) Score and reference coverage', loc='left', fontsize=FS)
    ax.set_xlim(0.40, 1.36)
    ax.set_ylim(0.36, 1.12)
    _p015(ax, fs=FS)
    ax.legend(fontsize=FS, loc="lower right", frameon=True, framealpha=0.92,
              handletextpad=0.4, labelspacing=0.35, borderpad=0.35)

    fig.tight_layout(pad=0.6, w_pad=1.4, h_pad=1.0)
    _p017(fig, OUT_FIG / "01_output.png", dpi=300)

    export_results(df,rho_pool)
    print("Saved: output/01_output.png and output/01_output.xlsx")


if __name__ == "__main__":
    main()
