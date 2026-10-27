
import os
for _key in ("OMP_NUM_THREADS","MKL_NUM_THREADS","OPENBLAS_NUM_THREADS"): os.environ.setdefault(_key,"1")


NEW_PROBLEMS = ['sphere2', 'sphere6', 'rosenbrock2', 'rosenbrock6', 'ackley6']
PROBLEMS = ['branin', 'hartmann3', 'borehole', 'sphere2', 'sphere6', 'rosenbrock2', 'rosenbrock6', 'ackley6']


from dataclasses import dataclass
import numpy as np
from scipy.stats import qmc


def _p001(n, d, seed):
    if n < 1:
        return np.empty((0, d))
    return qmc.Sobol(d, scramble=True, seed=seed).random_base2(int(np.ceil(np.log2(n))))[:n]


@dataclass
class Problem:
    name: str
    discrepancy: float = 1.0
    stress: bool = False

    @property
    def dim(self):
        if self.name not in PROBLEMS:raise ValueError('Unsupported benchmark: '+self.name)
        return {"branin": 2, "hartmann3": 3, "borehole": 8, **{name: int(name[-1]) for name in NEW_PROBLEMS}}[self.name]

    def raw(self, x):
        x = np.atleast_2d(np.asarray(x, dtype=float))
        if x.shape[1] != self.dim or np.any((x < 0) | (x > 1)):
            raise ValueError("Expected points in the unit cube with correct dimension")
        if self.name in NEW_PROBLEMS:
            family = self.name[:-1]
            if family == "sphere":
                z = -5.12 + 10.24*x
                return np.sum(z*z, axis=1)
            if family == "rosenbrock":
                z = -2.048 + 4.096*x
                return np.sum(100*(z[:,1:]-z[:,:-1]**2)**2+(1-z[:,:-1])**2, axis=1)
            if family == "ackley":

                z = -5 + 10*x
                return -20*np.exp(-.2*np.sqrt(np.mean(z*z,axis=1)))-np.exp(np.mean(np.cos(2*np.pi*z),axis=1))+20+np.e
        if self.name == "branin":
            a, b = 15*x[:, 0]-5, 15*x[:, 1]
            return (b-5.1*a*a/(4*np.pi**2)+5*a/np.pi-6)**2 + 10*(1-1/(8*np.pi))*np.cos(a)+10
        if self.name == "hartmann3":
            A = np.array([[3,10,30],[.1,10,35],[3,10,30],[.1,10,35]])
            P = 1e-4*np.array([[3689,1170,2673],[4699,4387,7470],[1091,8732,5547],[381,5743,8828]])
            return -np.sum(np.array([1,1.2,3,3.2])*np.exp(-np.sum(A*(x[:,None,:]-P)**2, axis=2)), axis=1)
        lo = np.array([.05,100,63070,990,63.1,700,1120,9855])
        hi = np.array([.15,50000,115600,1110,116,820,1680,12045])
        rw, r, tu, hu, tl, hl, length, kw = (lo+x*(hi-lo)).T
        lr = np.log(r/rw)
        return 2*np.pi*tu*(hu-hl)/(lr*(1+2*length*tu/(lr*rw**2*kw)+tu/tl))

    def base(self, x):

        if self.name in NEW_PROBLEMS:

            center, scale = {"sphere": (8*self.dim,8*self.dim),
                             "rosenbrock": (200*(self.dim-1),200*(self.dim-1)),
                             "ackley": (10.,10.)}[self.name[:-1]]
            return (self.raw(x)-center)/scale
        center, scale = {"branin": (50,50), "hartmann3": (-1,1),
                         "borehole": (80,50)}[self.name]
        return (self.raw(x)-center)/scale

    def delta(self, x, chain=1):
        x = np.atleast_2d(x)
        if self.stress:
            return self.discrepancy*(.7*np.sin(10*np.pi*x[:,0]) + .6*(x[:,1] > .5))
        if chain == 1:
            return self.discrepancy*(.18*np.sin(np.pi*x[:,0])+.12*(x[:,1]-.5)+.08*np.mean(x-.5, axis=1))
        return self.discrepancy*(-.12*np.cos(np.pi*x[:,1])+.10*x[:,0]*x[:,-1])

    def evaluate(self, x, level=2):
        if level not in (0, 1, 2):
            raise ValueError("level must be 0, 1, or 2")
        y = self.base(x)
        for chain in range(1, level+1):
            y = y+self.delta(x, chain)
        return y


import warnings
import numpy as np
from scipy.optimize import minimize
from sklearn.exceptions import ConvergenceWarning
from sklearn.gaussian_process import GaussianProcessRegressor
from sklearn.gaussian_process.kernels import ConstantKernel, Matern

def _p002(obj, initial, bounds):
    result = minimize(obj,initial,method="L-BFGS-B",jac=True,bounds=bounds,
                      options={"maxiter":45,"ftol":1e-7})
    return result.x, result.fun

class AdaptiveGP:

    def __init__(self):
        self.kernel = None
        self.last_optimized = -100

    def fit(self,x,y):
        x,y = np.asarray(x),np.asarray(y)
        kernel = self.kernel
        if kernel is None:
            kernel = ConstantKernel(1.,(.01,100))*Matern(np.full(x.shape[1],.35),(.04,4),nu=2.5)
        optimize = len(x)-self.last_optimized >= 8
        self.model = GaussianProcessRegressor(kernel=kernel,alpha=1e-6,normalize_y=True,
                                               optimizer=_p002 if optimize else None)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore",ConvergenceWarning)
            self.model.fit(x,y)
        self.kernel = self.model.kernel_
        if optimize:
            self.last_optimized = len(x)
        self.cache = {}
        return self

    def predict(self,x):
        x = np.atleast_2d(np.asarray(x,dtype=float))
        key = (x.shape,x.tobytes())
        if key not in self.cache:
            m,s = self.model.predict(x,return_std=True)
            if len(self.cache)>=5:
                self.cache.pop(next(iter(self.cache)))
            self.cache[key] = (m,np.maximum(s,1e-8))
        return self.cache[key]


import numpy as np
SEARCH_METHODS=("no_sobol","no_elite","boundary_pso","extreme_pso","sobol")

def _p003(task,level,budget,seed,method="boundary_pso",population=20):
    if method not in SEARCH_METHODS:
        raise ValueError(method)
    rng=np.random.default_rng(seed)
    n=min(population,budget)
    use_sobol=method in ("no_elite","sobol")
    initial=_p001(n,task.dim,seed) if use_sobol else rng.random((n,task.dim))
    if method=="sobol":
        x=_p001(budget,task.dim,seed)
        return x,task.evaluate(x,level)
    x=initial; velocity=np.zeros_like(x)
    bx=x.copy();bf=np.full(n,np.inf);xs=[];ys=[]
    count=0
    while count<budget:
        m=min(n,budget-count);current=x[:m]
        y=task.evaluate(current,level);xs.append(current.copy());ys.append(y);count+=m
        f=y if method=="extreme_pso" else abs(y)
        ids=np.flatnonzero(f<bf[:m]);bx[ids]=current[ids];bf[ids]=f[ids]
        fraction=count/budget; best=bx[np.argmin(bf)]
        velocity=(.85-.45*fraction)*velocity+1.5*rng.random(x.shape)*(bx-x)
        if method=="no_sobol":
            k=max(2,int(n*(.5-.4*fraction)))
            elite=bx[np.argsort(bf)[:k]][rng.integers(k,size=n)]
            diversity=np.mean(np.linalg.norm(x-best,axis=1))/np.sqrt(task.dim)
            strength=min(1.,3*diversity)

            leader=(1-strength)*best+strength*elite
        else:
            leader=best
        velocity+=1.5*rng.random(x.shape)*(leader-x)
        x=np.clip(x+np.clip(velocity,-.2,.2),0,1)

        reset=rng.random(n)<.05
        x[reset]=rng.random((sum(reset),task.dim))
    return np.vstack(xs),np.concatenate(ys)


import numpy as np
from scipy.spatial.distance import cdist
from scipy.stats import norm
from scipy.linalg import solve_triangular


def _p004(model,x):

    gp=model.model
    length=np.asarray(gp.kernel_.k2.length_scale)
    amplitude=float(gp.kernel_.k1.constant_value)
    diff=x[:,None,:]-gp.X_train_[None,:,:]
    radius=np.sqrt(np.sum((diff/length)**2,axis=2))
    factor=-amplitude*(5/3)*(1+np.sqrt(5)*radius)*np.exp(-np.sqrt(5)*radius)
    gradient=np.einsum('ns,nsd,s->nd',factor,diff/length**2,gp.alpha_)
    return gradient*float(np.asarray(gp._y_train_std).ravel()[0])


def _p005(model):

    gp=model.model
    inverse=solve_triangular(gp.L_,np.eye(len(gp.alpha_)),lower=True)
    diagonal=np.sum(inverse**2,axis=0)
    residual=gp.alpha_/diagonal*float(np.asarray(gp._y_train_std).ravel()[0])
    return float(np.mean(abs(residual)))


def _p006(model,particles,steps=5,targets=None):


    z=particles.copy()
    target=0. if targets is None else np.asarray(targets)
    for _ in range(steps):
        mu=model.model.predict(z);gradient=_p004(model,z)
        move=(mu-target)[:,None]*gradient/(np.sum(gradient**2,axis=1)[:,None]+1e-9)
        z=np.clip(z-np.clip(move,-.15,.15),0,1)
    return z


def _p007(model,particles,rng,swarm_steps=3):


    z=particles.copy();velocity=np.zeros_like(z)
    best=z.copy();fitness=abs(model.model.predict(z))
    for t in range(swarm_steps):
        current=abs(model.model.predict(z));better=current<fitness
        best[better]=z[better];fitness[better]=current[better]
        k=max(8,int(len(z)*(.20-.12*t/max(swarm_steps,1))))
        elite=best[np.argsort(fitness)[:k]]
        leaders=elite[rng.integers(len(elite),size=len(z))]
        velocity=.4*velocity+1.3*rng.random(z.shape)*(best-z)+.4*rng.random(z.shape)*(leaders-z)
        z=np.clip(z+np.clip(velocity,-.12,.12),0,1)
    return z


def _p008(task,level,budget,seed,config,return_trace=False):
    setting=config.get("initial",32)
    if setting=="dimension_adaptive":
        setting=min(48,max(24,8*task.dim))
    init=min(setting,budget)
    x=_p001(init,task.dim,seed)
    y=task.evaluate(x,level)
    model=AdaptiveGP();rng=np.random.default_rng(seed+310000)
    trace=[];iteration=0;exploration_guard=False
    while len(y)<budget:
        model.fit(x,y)
        guard=config.get("novelty_uncertainty",0.)
        error=_p005(model) if guard else 0.
        if config.get("adaptive_exploration",False):
            if error>config.get("novelty_threshold",.15):
                exploration_guard=True
            elif error<config.get("novelty_release",.05):
                exploration_guard=False
            novelty=guard if exploration_guard else 0.
        else:
            novelty=guard if guard and error>config.get("novelty_threshold",.15) else 0.


        population=_p001(config.get("population",1024),task.dim,seed+100003+iteration*31)
        archive=x[abs(y)<=.15]
        warming=len(y)<config.get("warmup",0)
        if len(archive) and not warming:
            ids=rng.integers(len(archive),size=len(population)//3)
            offspring=np.clip(archive[ids]+rng.normal(0,.12,size=(len(ids),task.dim)),0,1)
            population=np.vstack([population,offspring])
        raw_population=population.copy()
        if config.get("swarm_steps",0):
            population=_p007(model,population,rng,config["swarm_steps"])
        width=config.get("band_width",0.)
        targets=rng.uniform(-width,width,size=len(population)) if width else None
        candidates=_p006(model,population,config.get("projection_steps",5),targets)
        anchors=None;uncovered=None;influence=None;weights=None
        if config.get("coverage_gain",False):
            cloud=_p001(2048,task.dim,seed+200003)
            cm,cs=model.predict(cloud)
            cp=norm.cdf((.15-cm)/cs)-norm.cdf((-.15-cm)/cs)
            selected=cp>.25
            anchors=cloud[selected];weights=cp[selected]
            if len(anchors):
                anchor_distance=cdist(anchors,archive).min(axis=1) if len(archive) else np.ones(len(anchors))*np.sqrt(task.dim)
                uncovered=anchor_distance>.10*np.sqrt(task.dim)


                chosen=np.argsort(weights*(anchor_distance+.01))[-512:]
                candidates=np.vstack([candidates,anchors[chosen]])
                influence=cdist(candidates,anchors)<=.10*np.sqrt(task.dim)
        mu,sigma=model.predict(candidates)
        probability=norm.cdf((.15-mu)/sigma)-norm.cdf((-.15-mu)/sigma)


        separation=cdist(candidates,archive).min(axis=1) if len(archive) else np.ones(len(candidates))*np.sqrt(task.dim)
        observed_distance=cdist(candidates,x).min(axis=1)
        batch=[]
        size=min(config.get("batch",8),budget-len(y))
        explore=size if warming else config.get("exploration_slots",0)
        if not warming and config.get("adaptive_exploration",False) and not exploration_guard:
            explore=config.get("min_exploration_slots",1)
        for j in range(size):
            if explore and j>=size-explore:
                pm,ps=model.predict(raw_population)
                all_x=np.vstack([x,np.asarray(batch)]) if batch else x
                raw_distance=cdist(raw_population,all_x).min(axis=1)
                if novelty:
                    ps=np.sqrt(ps**2+(novelty*raw_distance)**2)
                straddle=ps*np.exp(-.5*(pm/ps)**2)*np.minimum(raw_distance/.1,1.)
                straddle[raw_distance<=1e-6]=-np.inf
                point=raw_population[int(np.argmax(straddle))].copy()
                batch.append(point)
                distance=np.linalg.norm(candidates-point,axis=1)
                separation=np.minimum(separation,distance)
                observed_distance=np.minimum(observed_distance,distance)
                if influence is not None:
                    uncovered &= np.linalg.norm(anchors-point,axis=1)>.10*np.sqrt(task.dim)
                continue
            score=np.maximum(probability,1e-8)**config.get("probability_power",1.)
            if influence is not None:
                gain=influence@(weights*uncovered)
                score*=gain+.01*np.maximum(separation,1e-5)
            else:
                score*=np.maximum(separation,1e-5)**config.get("diversity_power",1.)
            score[observed_distance<=1e-6]=-np.inf
            index=int(np.argmax(score))
            if not np.isfinite(score[index]):
                point=_p001(1,task.dim,seed+iteration*101+j+790001)[0]
            else:
                point=candidates[index].copy()
            batch.append(point)
            distance=np.linalg.norm(candidates-point,axis=1)
            separation=np.minimum(separation,distance)
            observed_distance=np.minimum(observed_distance,distance)
            if influence is not None:
                uncovered &= np.linalg.norm(anchors-point,axis=1)>.10*np.sqrt(task.dim)
        new_x=np.asarray(batch);new_y=task.evaluate(new_x,level)
        x=np.vstack([x,new_x]);y=np.append(y,new_y)
        trace.append(dict(queries=len(y),boundary_count=int(sum(abs(y)<=.15)),
                          kernel=str(model.kernel),loo_error=error,exploration_guard=exploration_guard))
        iteration+=1
    if return_trace:
        return x,y,trace
    return x,y


def _p009(h, a, b, z, sign):

    if z==1:
        c=b+90.;d=180.-c-a;b_angle=180.-2*d;a_angle=180.-b_angle+a-90.
    elif z==2 and a<b:
        c=90.-b;d=90.+a-b;b_angle=180.-2*d;a_angle=180.-b_angle-a-90.
    elif z==2 and a>b:
        c=90.-b;d=180.-c-a;b_angle=180.-2*d;a_angle=180.-b_angle-90.+a
    else:return 0.
    sd=np.sin(np.deg2rad(d));sa=np.sin(np.deg2rad(a_angle))
    with np.errstate(divide='ignore',invalid='ignore'):
        return float(np.sin(np.deg2rad(b_angle))*h*np.sin(np.deg2rad(c))/(sd*sa)*sign)

def _p010(hh, gamma, fai, omg):

    g,f,o=np.deg2rad(gamma),np.deg2rad(fai),np.deg2rad(omg)
    with np.errstate(divide='ignore',invalid='ignore',over='ignore'):
        return hh/np.tan(g)-((hh/np.sin(g)-hh*np.sin(f)/np.cos(f+g))*np.cos(o))/np.cos(o-g)

def _p011(fobj, nvars, lb, ub, budget, seed=0, population=20):


    if budget<2 or nvars<1 or population<2:raise ValueError('budget/population >= 2; nvars >= 1')
    lower=np.broadcast_to(np.asarray(lb,float),(nvars,));upper=np.broadcast_to(np.asarray(ub,float),(nvars,))
    if not np.all(upper>lower):raise ValueError('upper must exceed lower')
    rng=np.random.default_rng(seed);n=min(int(population),int(budget))
    pos=lower+rng.random((n,nvars))*(upper-lower)
    observations=[];objectives=[];curve=[];best=None;best_f=np.inf
    def evaluate(point):
        nonlocal best,best_f
        point=np.clip(np.asarray(point,float),lower,upper)


        if not np.isfinite(point).all():point=(best+lower+upper)/3 if best is not None else (lower+upper)/2
        value=float(fobj(point.copy()))
        if not np.isfinite(value):raise ValueError('Objective must be finite')
        observations.append(point.copy());objectives.append(value)
        if value<best_f:best_f=value;best=point.copy()
        curve.append(best_f)
        return point,value
    fitness=[]
    for i in range(n):pos[i],value=evaluate(pos[i]);fitness.append(value)
    fitness=np.asarray(fitness)
    def select(points,values):
        order=np.argsort(values,kind='stable')[:n]
        return points[order].copy(),values[order].copy()
    while len(observations)<budget:

        ac=rng.permutation(np.arange(1,n))
        q=len(observations)
        count=min(n-1,int(np.ceil(n*(2/3)*((budget-q+1)/budget))),budget-q)
        candidates=[];values=[]
        for index in ac[:count]:
            candidate=pos[index].copy()
            for k in range(nvars):
                h=(best[k]-pos[index,k])*rng.random()
                ceiling=5*np.arctanh(1-len(observations)/budget)+1
                h=max(1.,min(h,ceiling))
                sign=int(rng.integers(1,3))*2-3
                a=rng.random()*20;b=rng.random()*(45-a/2);z=int(rng.integers(1,3))
                candidate[k]+=_p009(h,a,b,z,sign)
            candidate,value=evaluate(candidate);candidates.append(candidate);values.append(value)
            if len(observations)==budget:break
        if candidates:pos,fitness=select(np.vstack([pos,candidates]),np.r_[fitness,values])
        if len(observations)==budget:break
        candidates=[];values=[]
        for j in range(n):
            hh=best-pos[j] if np.all(best!=pos[j]) else np.ones(nvars)*.05*(int(rng.integers(1,3))*2-3)
            direction=np.sign(hh);hh=np.abs(hh*rng.random(nvars))
            gamma=rng.random(nvars)*90*((budget-len(observations)*.99)/budget)
            amin=np.rad2deg(np.arctan(np.sin(np.deg2rad(gamma))*np.cos(np.deg2rad(gamma))/(1+np.sin(np.deg2rad(gamma))**2)))
            amax=np.rad2deg(np.arctan(1/(2*np.tan(np.deg2rad(gamma)))))
            fai=(amax-amin)*rng.random()+amin
            omg=np.rad2deg(np.arcsin(np.clip(rng.random()*np.sin(np.deg2rad(fai+gamma)),-1,1)))
            candidate,value=evaluate(pos[j]+_p010(hh,gamma,fai,omg)*direction)
            candidates.append(candidate);values.append(value)
            if len(observations)==budget:break
        pos,fitness=select(np.vstack([pos,candidates]),np.r_[fitness,values])
    return best,best_f,np.asarray(curve),np.asarray(observations)

def _p012(task,level,budget,seed,population=20):

    responses=[]
    def objective(point):
        value=float(task.evaluate(point[None,:],level)[0]);responses.append(value)
        return abs(value)
    _,_,_,points=_p011(objective,task.dim,0.,1.,budget,seed,population)
    return points,np.asarray(responses)


import argparse
from pathlib import Path
ROOT=Path(__file__).resolve().parent
SWARM_CONFIG={'initial': 24, 'population': 1024, 'batch': 8, 'projection_steps': 5, 'swarm_steps': 3, 'probability_power': 2.0, 'diversity_power': 1.0, 'band_width': 0.06, 'exploration_slots': 2, 'min_exploration_slots': 1, 'warmup': 0, 'coverage_gain': True, 'novelty_uncertainty': 2.0, 'novelty_threshold': 0.15, 'novelty_release': 0.075, 'adaptive_exploration': True}

def _p013(x,y,reference):
    b=np.unique(x[abs(y)<=.15],axis=0)
    reached=float(np.mean(cdist(reference,b).min(axis=1)/np.sqrt(x.shape[1])<=.10)) if len(b) and len(reference) else 0.
    return dict(unique_boundary_points=len(b),boundary_coverage=reached)


MSO_LICENSE = 'BSD 3-Clause License\n\nCopyright (c) 2025, 崔星星\n\nRedistribution and use in source and binary forms, with or without\nmodification, are permitted provided that the following conditions are met:\n\n1. Redistributions of source code must retain the above copyright notice, this\n   list of conditions and the following disclaimer.\n\n2. Redistributions in binary form must reproduce the above copyright notice,\n   this list of conditions and the following disclaimer in the documentation\n   and/or other materials provided with the distribution.\n\n3. Neither the name of the copyright holder nor the names of its\n   contributors may be used to endorse or promote products derived from\n   this software without specific prior written permission.\n\nTHIS SOFTWARE IS PROVIDED BY THE COPYRIGHT HOLDERS AND CONTRIBUTORS "AS IS"\nAND ANY EXPRESS OR IMPLIED WARRANTIES, INCLUDING, BUT NOT LIMITED TO, THE\nIMPLIED WARRANTIES OF MERCHANTABILITY AND FITNESS FOR A PARTICULAR PURPOSE ARE\nDISCLAIMED. IN NO EVENT SHALL THE COPYRIGHT HOLDER OR CONTRIBUTORS BE LIABLE\nFOR ANY DIRECT, INDIRECT, INCIDENTAL, SPECIAL, EXEMPLARY, OR CONSEQUENTIAL\nDAMAGES (INCLUDING, BUT NOT LIMITED TO, PROCUREMENT OF SUBSTITUTE GOODS OR\nSERVICES; LOSS OF USE, DATA, OR PROFITS; OR BUSINESS INTERRUPTION) HOWEVER\nCAUSED AND ON ANY THEORY OF LIABILITY, WHETHER IN CONTRACT, STRICT LIABILITY,\nOR TORT (INCLUDING NEGLIGENCE OR OTHERWISE) ARISING IN ANY WAY OUT OF THE USE\nOF THIS SOFTWARE, EVEN IF ADVISED OF THE POSSIBILITY OF SUCH DAMAGE.'


METHODS=('bapso',)+tuple(SEARCH_METHODS)+('mso',)
LABELS={'bapso':'bapso','sobol':'Sobol','boundary_pso':'Boundary PSO','mso':'MSO',
        'no_sobol':'Legacy no Sobol','no_elite':'Legacy no elite','extreme_pso':'Extreme PSO'}

def run_job(job):
    name,seed,budget=job;task=Problem(name)
    test=_p001(4096,task.dim,95000+seed)
    reference=test[abs(task.evaluate(test,0))<=.15]
    if not len(reference):raise ValueError('Empty evaluation boundary: '+name)
    rows=[]
    for method in METHODS:
        if method=='bapso':x,y=_p008(task,0,budget,seed,SWARM_CONFIG)
        elif method=='mso':x,y=_p012(task,0,budget,seed)
        else:x,y=_p003(task,0,budget,seed,method)
        assert len(y)==budget
        np.testing.assert_allclose(y,task.evaluate(x,0))
        if method=='bapso':assert len(np.unique(x,axis=0))==len(x)
        rows.append(dict(problem=name,seed=seed,method=method,**_p013(x,y,reference)))
    return rows


def _p014():
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,
        'axes.labelsize':9,'axes.titlesize':9,'xtick.labelsize':9,'ytick.labelsize':9,
        'legend.fontsize':9,'axes.edgecolor':'#333333','axes.linewidth':.8,
        'xtick.direction':'out','ytick.direction':'out','grid.color':'#D9D9D9',
        'grid.linewidth':.5,'grid.alpha':.6,'axes.axisbelow':True,
        'figure.facecolor':'white','axes.facecolor':'white','mathtext.fontset':'dejavusans'})

def plot_results(rows,problems,path,seeds=20,budget=160):
    summary=[]
    for name in problems:
        for method in METHODS:
            group=[r for r in rows if r['problem']==name and r['method']==method]
            count=np.array([r['unique_boundary_points'] for r in group]);cov=np.array([r['boundary_coverage'] for r in group])
            summary.append(dict(problem=name,method=method,count=count.mean(),sd_count=count.std(ddof=1) if len(count)>1 else 0,
                coverage=cov.mean(),sd_coverage=cov.std(ddof=1) if len(cov)>1 else 0))
    plot_summary(summary,problems,path)

def plot_summary(summary,problems,path):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    _p014()
    palette=['#2B5C8A','#C3D3DF','#7D9DB5','#5B8FC9','#D4A15A',
             '#B9C3CB','#B5443A']
    fig,axes=plt.subplots(len(problems),2,figsize=(16,2.8*len(problems)),squeeze=False)
    for i,name in enumerate(problems):
        group={r['method']:r for r in summary if r['problem']==name}
        for col,(key,sd,label) in enumerate([('count','sd_count','Boundary scenarios'),('coverage','sd_coverage','Coverage (%)')]):
            ax=axes[i,col];scale=100 if col else 1
            means=np.array([float(group[m][key])*scale for m in METHODS]);errors=np.array([float(group[m][sd])*scale for m in METHODS])
            bars=ax.bar(np.arange(len(METHODS)),means,yerr=errors,color=palette,edgecolor='#333333',linewidth=.4,capsize=2,error_kw={'elinewidth':.8})
            for bar,value,error in zip(bars,means,errors):
                ax.annotate(f'{value:.1f}',(bar.get_x()+bar.get_width()/2,bar.get_height()+error),xytext=(0,4),textcoords='offset points',ha='center',fontsize=9)
            ax.set_title(f'({chr(97+i*2+col)}) {name}',loc='left',pad=7)
            ax.set_ylabel(label);ax.set_xticks(range(len(METHODS)),[LABELS[m] for m in METHODS],rotation=32,ha='right')
            ax.set_ylim(0,max(means+errors)*1.3+.1);ax.grid(axis='y')
    fig.tight_layout(pad=1.2,h_pad=1.5,w_pad=1.8)
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);fig.savefig(path,dpi=300);plt.close(fig)


def _write_book(path,tables):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    import math,json
    workbook=Workbook(write_only=True)
    header_font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
    body_font=Font(name='Arial',size=10,color='222222')
    fill=PatternFill('solid',fgColor='2B5C8A')
    border=Border(bottom=Side(style='thin',color='CCCCCC'))
    header_alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    number_alignment=Alignment(horizontal='right',vertical='center')
    text_alignment=Alignment(horizontal='left',vertical='center')
    wrap_alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
    for title,records in tables:
        if not records:raise ValueError('Empty export table: '+title)
        sheet=workbook.create_sheet(title);headers=list(records[0])
        sheet.sheet_view.showGridLines=False
        sheet.freeze_panes='C2' if title in ('Trials','Curves','Observations') else 'A2'
        sheet.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{len(records)+1}'
        sheet.row_dimensions[1].height=42
        for column,key in enumerate(headers,1):
            sample=[r.get(key) for r in records[:100]]
            width=max(14,min(34,max([len(key.replace('_',' '))*.8]+[len(str(v)) for v in sample if isinstance(v,str)])))
            sheet.column_dimensions[get_column_letter(column)].width=90 if key=='value' else width
        cells=[]
        for value in headers:
            cell=WriteOnlyCell(sheet,value=value);cell.font=header_font;cell.fill=fill
            cell.alignment=header_alignment;cell.border=border;cells.append(cell)
        sheet.append(cells)
        for index,record in enumerate(records,2):
            cells=[]
            for key in headers:
                value=record.get(key)
                if isinstance(value,np.generic):value=value.item()
                if isinstance(value,(list,tuple,dict)):value=json.dumps(value,ensure_ascii=True)
                if isinstance(value,float) and not math.isfinite(value):raise ValueError('Non-finite export value')
                if isinstance(value,str) and not value.isascii():raise ValueError('Workbook text must be English ASCII')
                cell=WriteOnlyCell(sheet,value=value);cell.font=body_font
                number=isinstance(value,(int,float)) and not isinstance(value,bool)
                cell.alignment=number_alignment if number else text_alignment
                if key=='value' and isinstance(value,str) and len(value)>80:
                    cell.alignment=wrap_alignment;sheet.row_dimensions[index].height=15*int(np.ceil(len(value)/80))
                if isinstance(value,float):cell.number_format='0.00%' if 'coverage' in key or 'accuracy' in key else '0.00000'
                elif isinstance(value,int):cell.number_format='0'
                cells.append(cell)
            sheet.append(cells)
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize='9'
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:1'
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);workbook.save(path)

def _settings(values):
    import hashlib
    return [dict(parameter='source_script',value=Path(__file__).name),
            dict(parameter='source_sha256',value=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())]+[
            dict(parameter=key,value=value) for key,value in values.items()]


def export_results(rows,problems,seeds=20,budget=160):
    expected={(problem,seed,method) for problem in problems for seed in range(2000,2000+seeds) for method in METHODS}
    observed={(row['problem'],row['seed'],row['method']) for row in rows}
    if len(observed)!=len(rows) or observed!=expected:
        raise ValueError('Export rows must match the configured methods, problems and seeds exactly')
    summary=[]
    for problem in problems:
        local=[]
        for method in METHODS:
            group=[r for r in rows if r['problem']==problem and r['method']==method]
            count=np.array([r['unique_boundary_points'] for r in group],float)
            coverage=np.array([r['boundary_coverage'] for r in group],float)
            local.append(dict(problem=problem,method=method,repeats=len(group),
                boundary_count_mean=float(count.mean()),boundary_count_sd=float(count.std(ddof=1)) if len(count)>1 else 0.,
                coverage_mean=float(coverage.mean()),coverage_sd=float(coverage.std(ddof=1)) if len(coverage)>1 else 0.))
        for record in local:
            record['count_rank']=1+sum(r['boundary_count_mean']>record['boundary_count_mean']+1e-12 for r in local)
            record['coverage_rank']=1+sum(r['coverage_mean']>record['coverage_mean']+1e-12 for r in local)
        summary.extend(local)
    trials=[dict(problem=r['problem'],seed=r['seed'],method=r['method'],
        boundary_count=r['unique_boundary_points'],coverage=r['boundary_coverage']) for r in rows]
    settings=_settings(dict(problems=problems,methods=list(METHODS),seeds=list(range(2000,2000+seeds)),
        true_query_budget=budget,boundary_band=.15,normalized_coverage_distance=.10,evaluation_points=4096,
        metric_direction='Higher boundary count and coverage are better'))
    _write_book(ROOT/'output'/'02_output.xlsx',[('Summary',summary),('Trials',trials),('Settings',settings)])

def main():
    from concurrent.futures import ProcessPoolExecutor
    p=argparse.ArgumentParser(description='Single-platform testing; saves PNG and XLSX results')
    p.add_argument('--problems',nargs='+',choices=PROBLEMS,default=PROBLEMS)
    p.add_argument('--seeds',type=int,default=20);p.add_argument('--budget',type=int,default=160)
    p.add_argument('--workers',type=int,default=8);args=p.parse_args()
    if args.seeds<1 or args.budget<24 or args.workers<1:p.error('seeds>=1, budget>=24, workers>=1 required')
    jobs=[(name,seed,args.budget) for name in args.problems for seed in range(2000,2000+args.seeds)]
    rows=[]
    if args.workers==1:
        for job in jobs:rows.extend(run_job(job));print(job[:2],'complete',flush=True)
    else:
        with ProcessPoolExecutor(max_workers=args.workers) as pool:
            for i,result in enumerate(pool.map(run_job,jobs),1):rows.extend(result);print(f'Completed {i}/{len(jobs)}',flush=True)
    plot_results(rows,args.problems,ROOT/'output'/'02_output.png',args.seeds,args.budget)
    export_results(rows,args.problems,args.seeds,args.budget)
    print('Saved PNG and XLSX:',ROOT/'output')

if __name__=='__main__':main()
